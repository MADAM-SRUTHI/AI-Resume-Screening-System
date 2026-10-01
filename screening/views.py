from django.shortcuts import render, redirect
from .models import Resume
from pypdf import PdfReader
from docx import Document
from io import BytesIO
import re


def upload_resume(request):

    if request.method == "POST":

        name = request.POST.get("name")
        email = request.POST.get("email")
        resume_file = request.FILES.get("resume_file")

        if name and email and resume_file:

            # Read uploaded file
            file_data = resume_file.read()

            resume_text = ""

            # PDF file
            if resume_file.name.lower().endswith(".pdf"):

                reader = PdfReader(BytesIO(file_data))

                for page in reader.pages:
                    resume_text += page.extract_text() or ""

            # DOCX file
            elif resume_file.name.lower().endswith(".docx"):

                document = Document(BytesIO(file_data))

                # Read normal paragraphs
                for paragraph in document.paragraphs:
                    resume_text += paragraph.text + " "

                # Read tables also
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            resume_text += cell.text + " "

            # Skills list
            skills_list = [
                "Python",
                "Django",
                "Java",
                "C",
                "C++",
                "SQL",
                "HTML",
                "CSS",
                "JavaScript",
                "React",
                "Machine Learning",
                "Data Science"
            ]

            # Find skills
            found_skills = []

            for skill in skills_list:
                if re.search(
                    r"\b" + re.escape(skill) + r"\b",
                    resume_text,
                    re.IGNORECASE
                ):
                    found_skills.append(skill)

            skills = ", ".join(found_skills)

            # Reset file pointer
            resume_file.seek(0)

            # Save resume
            Resume.objects.create(
                name=name,
                email=email,
                resume_file=resume_file,
                skills=skills
            )

            return redirect("success")

    return render(request, "upload.html")


def success(request):
    return render(request, "success.html")